import base64
import tempfile
from datetime import datetime, timezone as datetime_timezone
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from django.test import override_settings
from django.utils import timezone
from docx import Document as DocxDocument
from PIL import Image
from cryptography.fernet import Fernet

from accounts.models import User
from ml_engine.face_recognition import FaceImageError, decrypt_embedding
from .forms import InternProfileForm
from .models import Application, AttendanceFaceCapture, AttendanceLog, CompanyProfile, DailyReport, FACE_CONSENT_VERSION, FaceEnrollment, InternProfile, Posting, RiskAssessment, Scorecard, WeeklyReport, WorksiteLocation
from .utils import haversine_distance_km
from .validators import FileSizeAndTypeValidator


class PortalWorkflowTests(TestCase):
    def setUp(self):
        self.intern_user = User.objects.create_user(username="intern", password="Safe-example-Password-927!")
        self.intern = InternProfile.objects.get(user=self.intern_user)
        self.company_user = User.objects.create_user(
            username="company",
            password="Safe-example-Password-927!",
            role=User.Role.COMPANY,
            first_name="Host",
            last_name="Company",
        )
        self.company = CompanyProfile.objects.get(user=self.company_user)
        self.posting = Posting.objects.create(
            company=self.company,
            title="Operations Intern",
            description="Support the placement team.",
            location="Remote",
            status=Posting.Status.PUBLISHED,
        )

    def _camera_data_url(self):
        image = BytesIO()
        Image.new("RGB", (100, 100), "white").save(image, format="JPEG")
        return "data:image/jpeg;base64," + base64.b64encode(image.getvalue()).decode("ascii")

    def _complete_intern_details(self, *, external=False):
        self.intern.student_id = "TEST-STUDENT-001"
        self.intern.university = "Example University"
        self.intern.course = "Information Technology"
        self.intern.year_level = 3
        self.intern.placement_type = (
            InternProfile.PlacementType.EXTERNAL if external else InternProfile.PlacementType.PLATFORM
        )
        self.intern.external_host = "External Host" if external else ""
        self.intern.save()

    def _complete_intern_profile(self, *, external=False):
        self._complete_intern_details(external=external)
        FaceEnrollment.objects.update_or_create(
            intern=self.intern,
            defaults={
                "encrypted_embedding": b"encrypted test embedding",
                "enrolled_by": self.intern_user,
                "consent_confirmed_at": timezone.now(),
                "consent_text_version": FACE_CONSENT_VERSION,
            },
        )
    def _enroll_attendance_location(self, intern=None):
        intern = intern or self.intern
        if intern.placement_type == InternProfile.PlacementType.EXTERNAL:
            return WorksiteLocation.objects.create(
                intern=intern,
                name="External training site",
                address="1 Training Road",
                latitude="14.600000",
                longitude="121.000000",
                created_by=self.company_user,
            )
        application, _ = Application.objects.get_or_create(
            intern=intern,
            posting=self.posting,
            defaults={"status": Application.Status.ACCEPTED},
        )
        if application.status != Application.Status.ACCEPTED:
            application.status = Application.Status.ACCEPTED
            application.save(update_fields=("status",))
        return WorksiteLocation.objects.create(
            company=self.company,
            name="Main office",
            address="1 Main Road",
            latitude="14.600000",
            longitude="121.000000",
            created_by=self.company_user,
        )

    def _attendance_payload(self, action, **extra):
        return {
            "action": action,
            f"{action}_latitude": "14.600000",
            f"{action}_longitude": "121.000000",
            **extra,
        }

    def test_company_dashboard_focuses_on_attendance_and_worksites(self):
        Application.objects.create(intern=self.intern, posting=self.posting)
        self.client.force_login(self.company_user)

        response = self.client.get(reverse("portal:company_dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.company.organization)
        self.assertContains(response, "Review attendance")
        self.assertContains(response, "Set up worksite locations")
        self.assertEqual(response.context["active_worksite_count"], 0)


    def test_company_can_enroll_and_deactivate_multiple_worksites(self):
        self.client.force_login(self.company_user)
        url = reverse("portal:company_worksites")
        payload = {
            "action": "create",
            "name": "North office",
            "address": "10 North Road",
            "latitude": "14.600000",
            "longitude": "121.000000",
            "radius_meters": "100",
        }

        created = self.client.post(url, payload)

        self.assertRedirects(created, url)
        worksite = WorksiteLocation.objects.get(company=self.company)
        self.assertEqual(worksite.name, "North office")
        self.assertEqual(worksite.radius_meters, 100)
        self.assertTrue(worksite.is_active)

        deactivated = self.client.post(url, {
            "action": "deactivate",
            "worksite_id": worksite.pk,
        })

        self.assertRedirects(deactivated, url)
        worksite.refresh_from_db()
        self.assertFalse(worksite.is_active)

    def test_coordinator_can_assign_external_intern_worksite(self):
        self._complete_intern_details(external=True)
        coordinator = User.objects.create_user(username="worksite-coordinator", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)
        url = reverse("portal:coordinator_external_worksites")

        response = self.client.post(url, {
            "action": "create",
            "intern": self.intern.pk,
            "name": "External host office",
            "address": "20 Host Road",
            "latitude": "14.600000",
            "longitude": "121.000000",
            "radius_meters": "100",
        })

        self.assertRedirects(response, url)
        worksite = WorksiteLocation.objects.get(intern=self.intern)
        self.assertIsNone(worksite.company_id)
        self.assertEqual(worksite.created_by, coordinator)

    def test_role_dashboard_routes_use_role_specific_paths(self):
        self.assertEqual(reverse("portal:intern_dashboard"), "/intern/dashboard/")
        self.assertEqual(reverse("portal:company_dashboard"), "/supervisor/dashboard/")
        self.assertEqual(reverse("portal:coordinator_dashboard"), "/admin/dashboard/")

    def test_current_navigation_item_is_highlighted(self):
        self._complete_intern_profile()
        self.client.force_login(self.intern_user)

        response = self.client.get(reverse("portal:intern_dashboard"))

        self.assertContains(response, 'class="nav-link-active" aria-current="page">Overview</a>')

    def test_intern_can_update_external_host_details_later(self):
        form = InternProfileForm(
            {
                "placement_type": InternProfile.PlacementType.EXTERNAL,
                "external_host": "Northside Design Studio",
                "student_id": "STUDENT-201",
                "university": "Example University",
                "course": "Design",
                "year_level": 4,
                "bio": "",
            },
            instance=self.intern,
        )

        self.assertTrue(form.is_valid(), form.errors)
        profile = form.save()
        self.assertEqual(profile.placement_type, InternProfile.PlacementType.EXTERNAL)
        self.assertEqual(profile.external_host, "Northside Design Studio")

        draft_form = InternProfileForm(
            {
                "placement_type": InternProfile.PlacementType.EXTERNAL,
                "external_host": "",
                "student_id": "STUDENT-202",
                "university": "Example University",
                "course": "Design",
                "year_level": 4,
                "bio": "",
            },
            instance=profile,
        )
        self.assertTrue(draft_form.is_valid(), draft_form.errors)
        draft_profile = draft_form.save()
        self.assertEqual(draft_profile.external_host, "")

    def test_intern_can_save_incomplete_profile_without_requirement_gate(self):
        self.client.force_login(self.intern_user)
        response = self.client.post(reverse("portal:profile"), {
            "student_id": "",
            "university": "",
            "course": "",
            "year_level": "",
            "placement_type": InternProfile.PlacementType.PLATFORM,
            "external_host": "",
            "bio": "",
        })

        self.assertRedirects(response, reverse("portal:profile"))
        self.intern.refresh_from_db()
        self.assertFalse(self.intern.student_id)
        self.assertFalse(self.intern.university)
        self.assertFalse(self.intern.course)
        self.assertIsNone(self.intern.year_level)
        self.assertEqual(self.client.get(reverse("portal:attendance")).status_code, 200)

    def test_intern_can_save_complete_profile_from_profile_page(self):
        self.client.force_login(self.intern_user)

        response = self.client.post(reverse("portal:profile"), {
            "student_id": "TEST-STUDENT-SAVE",
            "university": "Example University",
            "course": "Information Technology",
            "year_level": 3,
            "placement_type": InternProfile.PlacementType.PLATFORM,
            "external_host": "",
            "bio": "",
        })

        self.assertRedirects(response, reverse("portal:profile"))
        self.intern.refresh_from_db()
        self.assertEqual(self.intern.student_id, "TEST-STUDENT-SAVE")
        self.assertEqual(self.intern.university, "Example University")
        self.assertEqual(self.intern.course, "Information Technology")
        self.assertEqual(self.intern.year_level, 3)

    def test_profile_save_failure_is_logged_and_rolled_back(self):
        self.client.force_login(self.intern_user)

        with patch("portal.views.audit", side_effect=RuntimeError("audit unavailable")):
            with self.assertLogs("portal.views", level="ERROR") as captured_logs:
                with self.assertRaises(RuntimeError):
                    self.client.post(reverse("portal:profile"), {
                        "student_id": "TEST-STUDENT-ROLLBACK",
                        "university": "Example University",
                        "course": "Information Technology",
                        "year_level": 3,
                        "placement_type": InternProfile.PlacementType.PLATFORM,
                        "external_host": "",
                        "bio": "",
                    })

        self.assertIn("Profile update failed for intern user", captured_logs.output[0])
        self.intern.refresh_from_db()
        self.assertIsNone(self.intern.student_id)

    def test_coordinator_dashboard_shows_attendance_and_external_worksite_status(self):
        self.intern.placement_type = InternProfile.PlacementType.EXTERNAL
        self.intern.external_host = "Northside Design Studio"
        self.intern.save(update_fields=("placement_type", "external_host"))
        clock_in = datetime(2026, 10, 5, 8, tzinfo=datetime_timezone.utc)
        AttendanceLog.objects.create(
            intern=self.intern,
            clock_in=clock_in,
            clock_out=clock_in.replace(hour=16),
            time_in_approved=True,
            time_out_approved=True,
        )
        coordinator = User.objects.create_user(username="coordinator", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)

        response = self.client.get(reverse("portal:coordinator_dashboard"))

        self.assertContains(response, "External interns without a worksite")
        self.assertContains(response, "Assign external worksite locations")
        self.assertContains(response, "Daily reports to review")

    def test_external_intern_daily_report_is_reviewed_by_coordinator(self):
        self._complete_intern_profile(external=True)
        self.intern.external_host = "Northside Design Studio"
        self.intern.save(update_fields=("external_host",))
        self.client.force_login(self.intern_user)

        response = self.client.post(
            reverse("portal:daily_report"),
            {
                "accomplishments": "Prepared three design concepts.",
                "challenges": "Waiting for access to the shared drive.",
            },
        )

        self.assertRedirects(response, reverse("portal:daily_report"))
        report = DailyReport.objects.get(intern=self.intern)
        self.assertEqual(report.status, DailyReport.Status.SUBMITTED)

        coordinator = User.objects.create_user(username="coordinator", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)
        queue = self.client.get(reverse("portal:coordinator_daily_reports"))
        self.assertContains(queue, "Northside Design Studio")
        reviewed = self.client.post(
            reverse("portal:review_daily_report", args=(report.pk,)),
            {"decision": "review", "supervisor_feedback": "Good progress."},
        )

        self.assertRedirects(reviewed, reverse("portal:coordinator_daily_reports"))
        report.refresh_from_db()
        self.assertEqual(report.status, DailyReport.Status.REVIEWED)
        self.assertEqual(report.supervisor_feedback, "Good progress.")
        self.assertEqual(report.reviewer, coordinator)

    def test_platform_intern_can_submit_daily_report(self):
        self._complete_intern_profile()
        self.client.force_login(self.intern_user)

        response = self.client.post(
            reverse("portal:daily_report"),
            {
                "accomplishments": "Completed the assigned support tasks.",
                "challenges": "",
            },
        )

        self.assertRedirects(response, reverse("portal:daily_report"))
        report = DailyReport.objects.get(intern=self.intern)
        self.assertEqual(report.status, DailyReport.Status.SUBMITTED)

    def test_external_intern_uses_coordinator_workflow_instead_of_company_opportunities(self):
        self._complete_intern_profile(external=True)
        self.intern.external_host = "Northside Design Studio"
        self.intern.save(update_fields=("external_host",))
        self.client.force_login(self.intern_user)

        page = self.client.get(reverse("portal:postings"))
        self.assertContains(page, "Your external placement is managed directly with the OJT coordinator.")
        self.assertNotContains(page, self.posting.title)
        self.assertNotContains(page, ">Opportunities</a>")

        response = self.client.post(reverse("portal:apply", args=(self.posting.pk,)))

        self.assertRedirects(response, reverse("portal:postings"))
        self.assertFalse(Application.objects.filter(intern=self.intern).exists())


    def test_external_intern_attendance_appears_in_coordinator_dtr_queue(self):
        self.intern.placement_type = InternProfile.PlacementType.EXTERNAL
        self.intern.external_host = "Northside Design Studio"
        self.intern.save(update_fields=("placement_type", "external_host"))
        log = AttendanceLog.objects.create(intern=self.intern, work_date="2026-10-07", clock_in="2026-10-07T09:00:00Z")
        coordinator = User.objects.create_user(username="coordinator", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)

        response = self.client.get(reverse("portal:dtr_queue"))

        self.assertContains(response, "External: Northside Design Studio")
        self.assertContains(response, "Approve time in")
        approved = self.client.post(
            reverse("portal:approve_dtr", args=(log.pk,)),
            {"event": "time_in"},
        )
        self.assertRedirects(approved, reverse("portal:dtr_queue"))
        log.refresh_from_db()
        self.assertTrue(log.time_in_approved)

    def test_company_dtr_queue_excludes_external_placements(self):
        self.intern.placement_type = InternProfile.PlacementType.EXTERNAL
        self.intern.external_host = "Northside Design Studio"
        self.intern.save(update_fields=("placement_type", "external_host"))
        Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            status=Application.Status.ACCEPTED,
        )
        AttendanceLog.objects.create(
            intern=self.intern,
            work_date="2026-10-07",
            clock_in="2026-10-07T09:00:00Z",
        )
        self.client.force_login(self.company_user)

        response = self.client.get(reverse("portal:company_dtr_queue"))

        self.assertEqual(response.context["logs"].count(), 0)

    def test_new_intern_does_not_receive_ojt_requirements(self):
        other_user = User.objects.create_user(username="new-intern")
        other_intern = InternProfile.objects.get(user=other_user)

        self.assertFalse(other_intern.ojt_requirements.exists())

    def test_legacy_pages_redirect_to_core_workflows(self):
        self.client.force_login(self.intern_user)
        intern_page = self.client.get(reverse("portal:ojt_requirements"))
        self.assertRedirects(intern_page, reverse("portal:daily_report"))
        weekly_page = self.client.get(reverse("portal:reports"))
        self.assertRedirects(weekly_page, reverse("portal:daily_report"))

        self.client.force_login(self.company_user)
        company_page = self.client.get(reverse("portal:company_requirements"))
        self.assertRedirects(company_page, reverse("portal:company_dashboard"))

        coordinator = User.objects.create_user(username="coordinator", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)
        coordinator_page = self.client.get(reverse("portal:coordinator_ojt_requirements"))
        self.assertRedirects(coordinator_page, reverse("portal:coordinator_dashboard"))
        weekly_queue = self.client.get(reverse("portal:coordinator_weekly_reports"))
        self.assertRedirects(weekly_queue, reverse("portal:coordinator_daily_reports"))

    def test_attendance_does_not_require_legacy_ojt_approval(self):
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)

        response = self.client.post(
            reverse("portal:attendance"),
            self._attendance_payload("clock_in"),
        )

        self.assertRedirects(response, reverse("portal:attendance"))
        self.assertTrue(AttendanceLog.objects.filter(intern=self.intern).exists())

    def test_attendance_requires_location_inside_an_assigned_worksite(self):
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)

        outside = self.client.post(
            reverse("portal:attendance"),
            self._attendance_payload(
                "clock_in",
                clock_in_latitude="14.700000",
            ),
        )
        self.assertRedirects(outside, reverse("portal:attendance"))
        self.assertFalse(AttendanceLog.objects.filter(intern=self.intern).exists())

        invalid_coordinates = self.client.post(
            reverse("portal:attendance"),
            self._attendance_payload("clock_in", clock_in_latitude="NaN"),
        )
        self.assertRedirects(invalid_coordinates, reverse("portal:attendance"))
        self.assertFalse(AttendanceLog.objects.filter(intern=self.intern).exists())

        with patch(
            "portal.views._attendance_face_check",
            return_value=(AttendanceLog.FaceCheckStatus.MATCHED, 0.82, b"encrypted frame"),
        ):
            inside = self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload("clock_in"),
            )

        self.assertRedirects(inside, reverse("portal:attendance"))
        log = AttendanceLog.objects.get(intern=self.intern)
        self.assertEqual(str(log.clock_in_latitude), "14.600000")
        self.assertEqual(str(log.clock_in_longitude), "121.000000")
        self.assertTrue(log.clock_in_within_location)

    def test_attendance_face_preview_classifies_capture_without_recording_attendance(self):
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)
        cases = (
            (AttendanceLog.FaceCheckStatus.MATCHED, "good"),
            (AttendanceLog.FaceCheckStatus.NOT_MATCHED, "needs_review"),
            (AttendanceLog.FaceCheckStatus.UNAVAILABLE, "poor"),
        )

        for face_status, expected_quality in cases:
            with self.subTest(face_status=face_status), patch(
                "portal.views._attendance_face_check",
                return_value=(face_status, 0.81, b"unused preview image"),
            ) as face_check:
                response = self.client.post(
                    reverse("portal:attendance_face_preview"),
                    {"action": "clock_in", "face_image": self._camera_data_url()},
                )

                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["quality"], expected_quality)
                self.assertIn("label", response.json())
                self.assertIn("message", response.json())
                self.assertEqual(response.json()["similarity"], 0.81)
                face_check.assert_called_once()
                self.assertFalse(face_check.call_args.kwargs["retain_image"])

        self.assertFalse(AttendanceLog.objects.filter(intern=self.intern).exists())

    def test_attendance_face_preview_does_not_require_profile_or_enrollment(self):
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)

        response = self.client.post(
            reverse("portal:attendance_face_preview"),
            {"action": "clock_in", "face_image": self._camera_data_url()},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["quality"], "poor")
        self.assertIsNone(response.json()["similarity"])

    def test_attendance_face_preview_returns_json_for_invalid_action(self):
        self._complete_intern_profile()
        self.client.force_login(self.intern_user)

        response = self.client.post(
            reverse("portal:attendance_face_preview"),
            {"action": "invalid", "face_image": self._camera_data_url()},
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json(), {"error": "Unknown attendance action."})

    def test_company_cannot_use_intern_face_preview_endpoint(self):
        self.client.force_login(self.company_user)

        response = self.client.post(
            reverse("portal:attendance_face_preview"),
            {"action": "clock_in", "face_image": self._camera_data_url()},
        )

        self.assertEqual(response.status_code, 403)

    @override_settings(FACE_EMBEDDING_ENCRYPTION_KEY=Fernet.generate_key().decode("ascii"))
    def test_intern_can_enroll_own_face_from_profile_after_consent(self):
        self.client.force_login(self.intern_user)
        profile_url = reverse("portal:profile")
        enrollment_url = reverse("portal:intern_face_enrollment")
        self.assertContains(self.client.get(profile_url), "Face enrollment for attendance")
        payload = {"face_image": self._camera_data_url()}

        missing_consent = self.client.post(enrollment_url, payload)
        self.assertRedirects(missing_consent, profile_url)
        self.assertFalse(FaceEnrollment.objects.filter(intern=self.intern).exists())

        payload["consent_confirmed"] = "on"
        with patch("portal.views.create_face_embedding", return_value=b"intern embedding"):
            enrolled = self.client.post(enrollment_url, payload)

        self.assertRedirects(enrolled, profile_url)
        record = FaceEnrollment.objects.get(intern=self.intern)
        self.assertEqual(record.enrolled_by, self.intern_user)
        self.assertEqual(record.consent_text_version, FACE_CONSENT_VERSION)
        self.assertEqual(decrypt_embedding(record.encrypted_embedding), b"intern embedding")
        self.assertContains(self.client.get(profile_url), "Last updated")

    @override_settings(FACE_EMBEDDING_ENCRYPTION_KEY=Fernet.generate_key().decode("ascii"))
    def test_intern_face_enrollment_preview_requires_detectable_single_face(self):
        self.client.force_login(self.intern_user)
        with patch("portal.views.create_face_embedding", return_value=b"preview embedding") as create_embedding:
            response = self.client.post(
                reverse("portal:face_enrollment_preview"),
                {"face_image": self._camera_data_url()},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["quality"], "good")
        self.assertTrue(response.json()["label"])
        create_embedding.assert_called_once()
        self.assertFalse(FaceEnrollment.objects.filter(intern=self.intern).exists())

        with patch(
            "portal.views.create_face_embedding",
            side_effect=FaceImageError("Capture must contain exactly one clearly visible face."),
        ):
            poor_capture = self.client.post(
                reverse("portal:face_enrollment_preview"),
                {"face_image": self._camera_data_url()},
            )
        self.assertEqual(poor_capture.status_code, 200)
        self.assertEqual(poor_capture.json()["quality"], "poor")

    def test_live_face_detection_preview_returns_face_count_without_profile_or_enrollment(self):
        self.client.force_login(self.intern_user)
        with patch(
            "portal.views.detect_faces",
            return_value={"face_detected": True, "face_count": 1, "detections": []},
        ) as detector:
            response = self.client.post(
                reverse("portal:face_detection_preview"),
                {"face_image": self._camera_data_url()},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"face_count": 1, "face_detected": True})
        detector.assert_called_once()

    @override_settings(FACE_CAPTURE_DIAGNOSTICS=True)
    def test_live_face_detection_logs_decoded_image_shape_when_no_face_is_found(self):
        self.client.force_login(self.intern_user)
        with (
            patch(
                "portal.views.detect_faces",
                return_value={"face_detected": False, "face_count": 0, "detections": []},
            ),
            self.assertLogs("portal.views", level="WARNING") as captured_logs,
        ):
            response = self.client.post(
                reverse("portal:face_detection_preview"),
                {"face_image": self._camera_data_url()},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"face_count": 0, "face_detected": False})
        self.assertIn("image_type=Image", captured_logs.output[0])
        self.assertIn("mode=RGB", captured_logs.output[0])
        self.assertIn("array_shape=(100, 100, 3)", captured_logs.output[0])
        self.assertIn("array_dtype=uint8", captured_logs.output[0])
        self.assertIn("array_empty=False", captured_logs.output[0])
        self.assertIn("face_count=0", captured_logs.output[0])

    def test_face_capture_api_reports_expired_session_as_json(self):
        response = self.client.post(
            reverse("portal:face_detection_preview"),
            {"face_image": self._camera_data_url()},
        )

        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json(), {"error": "Your session expired. Sign in again and retry."})

    def test_intern_dashboard_and_attendance_are_available_without_ojt_approval(self):
        self.client.force_login(self.intern_user)

        dashboard = self.client.get(reverse("portal:intern_dashboard"))
        attendance = self.client.get(reverse("portal:attendance"))
        self.assertEqual(dashboard.status_code, 200)
        self.assertEqual(attendance.status_code, 200)
        self.assertContains(attendance, "Worksite location is not set up")

    def test_attendance_page_has_face_guide_and_capture_decision_controls(self):
        self._complete_intern_profile()
        self.client.force_login(self.intern_user)

        response = self.client.get(reverse("portal:attendance"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "face-guide")
        self.assertContains(response, "face-detection-guide.js")
        self.assertContains(response, "Capture and check")
        self.assertContains(response, "Retake face")
        self.assertContains(response, "Submit for approval")

    def test_face_mismatch_does_not_record_attendance_and_creates_review_request(self):
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)

        with patch(
            "portal.views._attendance_face_check",
            return_value=(AttendanceLog.FaceCheckStatus.NOT_MATCHED, 0.21, b"encrypted test capture"),
        ):
            response = self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload("clock_in"),
            )

        self.assertRedirects(response, reverse("portal:attendance"))
        log = AttendanceLog.objects.get(intern=self.intern)
        self.assertIsNone(log.clock_in)
        self.assertEqual(str(log.clock_in_latitude), "14.600000")
        self.assertTrue(log.clock_in_within_location)
        self.assertEqual(log.clock_in_face_status, AttendanceLog.FaceCheckStatus.NOT_MATCHED)
        self.assertEqual(log.clock_in_face_score, 0.21)
        capture = log.face_captures.get(event=AttendanceFaceCapture.Event.TIME_IN)
        self.assertEqual(capture.review_status, AttendanceFaceCapture.ReviewStatus.PENDING)
        self.assertEqual(bytes(capture.encrypted_image), b"encrypted test capture")

    def test_rejecting_face_exception_preserves_the_location_verified_state(self):
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)

        with patch(
            "portal.views._attendance_face_check",
            return_value=(AttendanceLog.FaceCheckStatus.NOT_MATCHED, 0.21, b"encrypted test capture"),
        ):
            self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload("clock_in"),
            )

        log = AttendanceLog.objects.get(intern=self.intern)
        capture = log.face_captures.get(event=AttendanceFaceCapture.Event.TIME_IN)
        self.assertTrue(log.clock_in_within_location)
        self.assertEqual(str(log.clock_in_latitude), "14.600000")

        coordinator = User.objects.create_user(username="location-reviewer", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)
        rejected = self.client.post(
            reverse("portal:review_attendance_face_capture", args=(capture.pk,)),
            {"action": "reject"},
        )

        self.assertRedirects(rejected, reverse("portal:dtr_queue"))
        log.refresh_from_db()
        self.assertTrue(log.clock_in_within_location)
        self.assertEqual(str(log.clock_in_latitude), "14.600000")
        self.assertEqual(str(log.clock_in_longitude), "121.000000")

    def test_company_can_approve_no_camera_time_in_without_server_error(self):
        Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            status=Application.Status.ACCEPTED,
        )
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)
        submitted = self.client.post(
            reverse("portal:attendance"),
            self._attendance_payload("clock_in"),
        )
        self.assertRedirects(submitted, reverse("portal:attendance"))

        log = AttendanceLog.objects.get(intern=self.intern)
        capture = log.face_captures.get(event=AttendanceFaceCapture.Event.TIME_IN)
        self.assertEqual(capture.face_status, AttendanceLog.FaceCheckStatus.NOT_CAPTURED)
        self.assertIsNone(capture.encrypted_image)
        self.assertIsNone(log.clock_in)

        self.client.force_login(self.company_user)
        queue = self.client.get(reverse("portal:company_dtr_queue"))
        self.assertEqual(queue.status_code, 200)
        approved = self.client.post(
            reverse("portal:review_attendance_face_capture", args=(capture.pk,)),
            {"action": "approve"},
        )
        self.assertRedirects(approved, reverse("portal:company_dtr_queue"))

        log.refresh_from_db()
        capture.refresh_from_db()
        self.assertIsNotNone(log.clock_in)
        self.assertTrue(log.clock_in_within_location)
        self.assertTrue(log.time_in_approved)
        self.assertEqual(capture.review_status, AttendanceFaceCapture.ReviewStatus.APPROVED)

    def test_attendance_page_explains_required_face_verification(self):
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)

        response = self.client.get(reverse("portal:attendance"))

        self.assertContains(response, "Face verification required")
        self.assertContains(response, "retained for up to 30 days")
        self.assertContains(response, "Camera starts when you choose Time in or Time out.")
        self.assertNotContains(response, "Start camera")
        self.assertNotContains(response, "Capture face for attendance")
        self.assertContains(response, "Capture and check")
        self.assertContains(response, 'name="face_image"')

    def test_successful_face_match_waits_for_approval_before_each_clock_event(self):
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)
        with patch(
            "portal.views._attendance_face_check",
            side_effect=[
                (AttendanceLog.FaceCheckStatus.MATCHED, 0.81, b"encrypted time-in"),
                (AttendanceLog.FaceCheckStatus.MATCHED, 0.79, b"encrypted time-out"),
            ],
        ):
            time_in_response = self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload("clock_in"),
            )
            self.assertRedirects(time_in_response, reverse("portal:attendance"))
            log = AttendanceLog.objects.get(intern=self.intern)
            time_in_capture = log.face_captures.get(event=AttendanceFaceCapture.Event.TIME_IN)
            self.assertEqual(time_in_capture.review_status, AttendanceFaceCapture.ReviewStatus.PENDING)
            self.assertFalse(log.time_in_approved)

            blocked_time_out = self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload("clock_out"),
            )
            self.assertRedirects(blocked_time_out, reverse("portal:attendance"))
            log.refresh_from_db()
            self.assertIsNone(log.clock_out)

            coordinator = User.objects.create_user(username="attendance-approver", role=User.Role.COORDINATOR)
            self.client.force_login(coordinator)
            queue = self.client.get(reverse("portal:dtr_queue"))
            self.assertContains(queue, "Approve time in")
            self.assertEqual(queue.context["logs"][0].pending_time_in_face_capture, time_in_capture)
            bypass = self.client.post(
                reverse("portal:approve_dtr", args=(log.pk,)),
                {"event": "time_in"},
            )
            self.assertRedirects(bypass, reverse("portal:dtr_queue"))
            log.refresh_from_db()
            self.assertFalse(log.time_in_approved)
            approved_time_in = self.client.post(
                reverse("portal:review_attendance_face_capture", args=(time_in_capture.pk,)),
                {"action": "approve"},
            )
            self.assertRedirects(approved_time_in, reverse("portal:dtr_queue"))
            self.client.force_login(self.intern_user)
            time_out_response = self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload("clock_out"),
            )

        self.assertRedirects(time_out_response, reverse("portal:attendance"))
        log.refresh_from_db()
        time_out_capture = log.face_captures.get(event=AttendanceFaceCapture.Event.TIME_OUT)
        self.assertIsNotNone(log.clock_in)
        self.assertIsNotNone(log.clock_out)
        self.assertFalse(log.time_out_approved)
        self.assertEqual(time_out_capture.review_status, AttendanceFaceCapture.ReviewStatus.PENDING)

        self.client.force_login(coordinator)
        approved_time_out = self.client.post(
            reverse("portal:review_attendance_face_capture", args=(time_out_capture.pk,)),
            {"action": "approve"},
        )
        self.assertRedirects(approved_time_out, reverse("portal:dtr_queue"))
        log.refresh_from_db()
        self.assertTrue(log.time_in_approved)
        self.assertTrue(log.time_out_approved)
        self.assertTrue(log.clock_in_within_location)
        self.assertTrue(log.clock_out_within_location)
        self.assertEqual(log.face_captures.count(), 2)
        self.assertEqual(
            set(log.face_captures.values_list("review_status", flat=True)),
            {AttendanceFaceCapture.ReviewStatus.APPROVED},
        )

    @override_settings(FACE_EMBEDDING_ENCRYPTION_KEY=Fernet.generate_key().decode("ascii"))
    def test_company_can_review_failure_and_view_only_its_attendance_capture(self):
        from ml_engine.face_recognition import encrypt_face_capture

        Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            status=Application.Status.ACCEPTED,
        )
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)
        encrypted_image = encrypt_face_capture(b"test attendance jpeg")
        with patch(
            "portal.views._attendance_face_check",
            return_value=(AttendanceLog.FaceCheckStatus.NOT_MATCHED, 0.21, encrypted_image),
        ):
            self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload("clock_in"),
            )

        log = AttendanceLog.objects.get(intern=self.intern)
        capture = log.face_captures.get()
        self.assertIsNone(log.clock_in)
        self.client.force_login(self.company_user)
        queue = self.client.get(reverse("portal:company_dtr_queue"))
        self.assertContains(queue, "Approve time in")
        self.assertContains(queue, "Awaiting company/admin approval")

        image_response = self.client.get(
            reverse("portal:attendance_face_capture_image", args=(capture.pk,))
        )
        self.assertEqual(image_response.status_code, 200)
        self.assertEqual(image_response.content, b"test attendance jpeg")
        self.assertEqual(image_response["Cache-Control"], "private, no-store")

        other_user = User.objects.create_user(username="other-face-intern", role=User.Role.INTERN)
        other_intern, _ = InternProfile.objects.get_or_create(user=other_user)
        other_log = AttendanceLog.objects.create(intern=other_intern, work_date="2026-10-09")
        other_capture = AttendanceFaceCapture.objects.create(
            attendance_log=other_log,
            event=AttendanceFaceCapture.Event.TIME_IN,
            face_status=AttendanceLog.FaceCheckStatus.ERROR,
            encrypted_image=encrypted_image,
        )
        forbidden_image = self.client.get(
            reverse("portal:attendance_face_capture_image", args=(other_capture.pk,))
        )
        self.assertEqual(forbidden_image.status_code, 404)

        reviewed = self.client.post(
            reverse("portal:review_attendance_face_capture", args=(capture.pk,)),
            {"action": "approve"},
        )
        self.assertRedirects(reviewed, reverse("portal:company_dtr_queue"))
        log.refresh_from_db()
        capture.refresh_from_db()
        self.assertEqual(log.clock_in, capture.captured_at)
        self.assertTrue(log.time_in_approved)
        self.assertEqual(capture.review_status, AttendanceFaceCapture.ReviewStatus.APPROVED)
        coordinator = User.objects.create_user(username="capture-review-coordinator", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)
        coordinator_queue = self.client.get(reverse("portal:dtr_queue"))
        self.assertContains(coordinator_queue, "View capture")
        coordinator_image = self.client.get(
            reverse("portal:attendance_face_capture_image", args=(capture.pk,))
        )
        self.assertEqual(coordinator_image.status_code, 200)
        self.assertEqual(coordinator_image.content, b"test attendance jpeg")

    def test_expired_attendance_capture_image_is_deleted_by_retention_command(self):
        log = AttendanceLog.objects.create(intern=self.intern, work_date="2026-10-09")
        capture = AttendanceFaceCapture.objects.create(
            attendance_log=log,
            event=AttendanceFaceCapture.Event.TIME_IN,
            face_status=AttendanceLog.FaceCheckStatus.ERROR,
            encrypted_image=b"expired encrypted image",
            image_expires_at=datetime(2020, 1, 1, tzinfo=datetime_timezone.utc),
        )

        call_command("purge_face_attendance_images", verbosity=0)

        capture.refresh_from_db()
        self.assertIsNone(capture.encrypted_image)
        self.assertTrue(AttendanceFaceCapture.objects.filter(pk=capture.pk).exists())

    @override_settings(FACE_EMBEDDING_ENCRYPTION_KEY=Fernet.generate_key().decode("ascii"))
    def test_coordinator_enrollment_requires_consent_and_stores_only_encrypted_embedding(self):
        coordinator = User.objects.create_user(username="face-coordinator", role=User.Role.COORDINATOR)
        self.client.force_login(coordinator)
        url = reverse("portal:coordinator_face_enrollment")
        self.assertContains(self.client.get(url), "Face-recognition pilot enrollment")
        payload = {
            "action": "enroll",
            "intern_id": self.intern.pk,
            "face_image": self._camera_data_url(),
        }

        missing_consent = self.client.post(url, payload)
        self.assertRedirects(missing_consent, url)
        self.assertFalse(FaceEnrollment.objects.filter(intern=self.intern).exists())

        payload["consent_confirmed"] = "on"
        with patch("portal.views.create_face_embedding", return_value=b"test embedding"):
            enrolled = self.client.post(url, payload)

        self.assertRedirects(enrolled, url)
        record = FaceEnrollment.objects.get(intern=self.intern)
        self.assertNotEqual(bytes(record.encrypted_embedding), b"test embedding")
        self.assertEqual(decrypt_embedding(record.encrypted_embedding), b"test embedding")
        self.assertEqual(record.enrolled_by, coordinator)
        self.assertIsNotNone(record.consent_confirmed_at)
        AttendanceLog.objects.create(
            intern=self.intern,
            work_date="2026-10-09",
            clock_in_face_status=AttendanceLog.FaceCheckStatus.MATCHED,
            clock_in_face_score=0.8,
        )

        deleted = self.client.post(url, {
            "action": "delete",
            "intern_id": self.intern.pk,
        })
        self.assertRedirects(deleted, url)
        self.assertFalse(FaceEnrollment.objects.filter(intern=self.intern).exists())
        self.assertEqual(
            AttendanceLog.objects.get(intern=self.intern).clock_in_face_status,
            AttendanceLog.FaceCheckStatus.NOT_ATTEMPTED,
        )




    def test_intern_dashboard_shows_approved_hours_and_core_tasks(self):
        self._complete_intern_profile()
        clock_in = datetime(2026, 10, 5, 8, tzinfo=datetime_timezone.utc)
        AttendanceLog.objects.create(
            intern=self.intern,
            clock_in=clock_in,
            clock_out=clock_in.replace(hour=16),
            time_in_approved=True,
            time_out_approved=True,
        )
        self.client.force_login(self.intern_user)

        response = self.client.get(reverse("portal:intern_dashboard"))

        self.assertEqual(response.context["completed_hours"], 8)
        self.assertContains(response, "Approved hours")
        self.assertContains(response, "Daily reports")

    def test_intern_progress_requires_both_time_approvals(self):
        self._complete_intern_profile()
        clock_in = datetime(2026, 10, 5, 8, tzinfo=datetime_timezone.utc)
        log = AttendanceLog.objects.create(
            intern=self.intern,
            clock_in=clock_in,
            clock_out=clock_in.replace(hour=16),
            time_in_approved=True,
        )
        self.client.force_login(self.intern_user)

        response = self.client.get(reverse("portal:intern_dashboard"))
        self.assertEqual(response.context["completed_hours"], 0)

        log.time_out_approved = True
        log.save(update_fields=("time_out_approved",))
        response = self.client.get(reverse("portal:intern_dashboard"))
        self.assertEqual(response.context["completed_hours"], 8)

    def test_coordinator_dashboard_focuses_on_attendance_reports_and_worksites(self):
        coordinator = User.objects.create_user(
            username="coordinator",
            password="Safe-example-Password-927!",
            role=User.Role.COORDINATOR,
        )
        self.client.force_login(coordinator)

        response = self.client.get(reverse("portal:coordinator_dashboard"))

        self.assertContains(response, "Pending DTR events")
        self.assertContains(response, "Daily reports to review")
        self.assertNotContains(response, "Intern directory")

    def test_company_can_export_accepted_intern_attendance_csv(self):
        self.intern.student_id = "STU-CSV-1"
        self.intern.save(update_fields=("student_id",))
        Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            status=Application.Status.ACCEPTED,
        )
        clock_in = datetime(2026, 10, 5, 8, tzinfo=datetime_timezone.utc)
        AttendanceLog.objects.create(
            intern=self.intern,
            clock_in=clock_in,
            clock_out=clock_in.replace(hour=16),
            time_in_approved=True,
            time_out_approved=True,
        )
        self.client.force_login(self.company_user)

        response = self.client.get(reverse("portal:company_attendance_export"))

        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertContains(response, "STU-CSV-1")

    def test_company_profile_can_be_updated(self):
        self.client.force_login(self.company_user)

        response = self.client.post(
            reverse("portal:company_profile"),
            {
                "organization": "Updated Host Organization",
                "website": "https://example.org",
                "address": "12 Fieldwork Road",
            },
        )

        self.assertRedirects(response, reverse("portal:company_profile"))
        self.company.refresh_from_db()
        self.assertEqual(self.company.organization, "Updated Host Organization")
        self.assertEqual(self.company.website, "https://example.org")
        self.assertEqual(self.company.address, "12 Fieldwork Road")

    def test_applicant_screening_shows_pending_applications(self):
        self.intern.course = "Information Systems"
        self.intern.save(update_fields=("course",))
        application = Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            cover_letter="I am interested in this placement.",
        )
        self.client.force_login(self.company_user)

        response = self.client.get(reverse("portal:applicant_screening"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.intern_user.username)
        self.assertContains(response, "Information Systems")
        self.assertContains(response, "I am interested in this placement.")
        self.assertEqual(list(response.context["pending"]), [application])

    def test_company_can_preview_and_download_own_resume(self):
        pdf_content = b"%PDF-1.7\nprivate resume"
        self.client.force_login(self.company_user)

        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            application = Application.objects.create(
                intern=self.intern,
                posting=self.posting,
                resume=SimpleUploadedFile("candidate.pdf", pdf_content, content_type="application/pdf"),
            )

            preview = self.client.get(
                reverse("portal:company_application_resume", args=(application.pk, "preview"))
            )
            self.assertEqual(preview.status_code, 200)
            self.assertEqual(preview["Content-Type"], "application/pdf")
            self.assertIn("inline", preview["Content-Disposition"])
            self.assertEqual(b"".join(preview.streaming_content), pdf_content)
            self.assertEqual(preview["Cache-Control"], "private, no-store")

            download = self.client.get(
                reverse("portal:company_application_resume", args=(application.pk, "download"))
            )
            self.assertEqual(download.status_code, 200)
            self.assertIn("attachment", download["Content-Disposition"])
            self.assertEqual(b"".join(download.streaming_content), pdf_content)

    def test_company_can_preview_docx_resume_as_html(self):
        document_stream = BytesIO()
        document = DocxDocument()
        document.add_heading("Candidate Resume", level=1)
        document.add_paragraph("Experienced operations intern.")
        document.save(document_stream)
        self.client.force_login(self.company_user)

        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            application = Application.objects.create(
                intern=self.intern,
                posting=self.posting,
                resume=SimpleUploadedFile(
                    "candidate.docx",
                    document_stream.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            )

            applications_response = self.client.get(reverse("portal:company_applications"))
            screening_response = self.client.get(reverse("portal:applicant_screening"))
            self.assertContains(
                applications_response,
                reverse("portal:company_application_resume", args=(application.pk, "preview")),
            )
            self.assertContains(screening_response, "Preview resume")

            response = self.client.get(
                reverse("portal:company_application_resume", args=(application.pk, "preview"))
            )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "portal/company/resume_preview.html")
        self.assertContains(response, "Candidate Resume")
        self.assertContains(response, "Experienced operations intern.")
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_company_cannot_preview_another_company_resume(self):
        other_user = User.objects.create_user(
            username="resume-owner-company",
            password="Safe-example-Password-927!",
            role=User.Role.COMPANY,
        )
        other_company = CompanyProfile.objects.get(user=other_user)
        other_posting = Posting.objects.create(
            company=other_company,
            title="Private Placement",
            description="Private application.",
            location="Remote",
        )
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            application = Application.objects.create(
                intern=self.intern,
                posting=other_posting,
                resume=SimpleUploadedFile("candidate.pdf", b"%PDF-1.7\nprivate resume", content_type="application/pdf"),
            )
            self.client.force_login(self.company_user)

            response = self.client.get(
                reverse("portal:company_application_resume", args=(application.pk, "preview"))
            )

        self.assertEqual(response.status_code, 404)

    def test_company_can_accept_own_pending_application(self):
        application = Application.objects.create(intern=self.intern, posting=self.posting)
        self.client.force_login(self.company_user)

        response = self.client.post(
            reverse("portal:application_decide", args=(application.pk, "accept"))
        )

        self.assertRedirects(response, reverse("portal:applicant_screening"))
        application.refresh_from_db()
        self.assertEqual(application.status, Application.Status.ACCEPTED)

    def test_company_cannot_decide_another_company_application(self):
        other_user = User.objects.create_user(
            username="other-company",
            password="Safe-example-Password-927!",
            role=User.Role.COMPANY,
        )
        other_company = CompanyProfile.objects.get(user=other_user)
        other_posting = Posting.objects.create(
            company=other_company,
            title="Other Placement",
            description="A different placement.",
            location="Remote",
        )
        application = Application.objects.create(intern=self.intern, posting=other_posting)
        self.client.force_login(self.company_user)

        response = self.client.post(
            reverse("portal:application_decide", args=(application.pk, "accept"))
        )

        self.assertEqual(response.status_code, 404)
        application.refresh_from_db()
        self.assertEqual(application.status, Application.Status.SUBMITTED)

    def test_company_dtr_only_shows_accepted_interns_and_can_approve(self):
        Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            status=Application.Status.ACCEPTED,
        )
        own_log = AttendanceLog.objects.create(
            intern=self.intern,
            work_date="2026-10-01",
            clock_in="2026-10-01T09:00:00Z",
            clock_out="2026-10-01T17:00:00Z",
        )
        other_user = User.objects.create_user(username="other-intern", password="Safe-example-Password-927!")
        other_intern = InternProfile.objects.get(user=other_user)
        AttendanceLog.objects.create(intern=other_intern, work_date="2026-10-01")
        self.client.force_login(self.company_user)

        response = self.client.get(reverse("portal:company_dtr_queue"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.intern_user.username)
        self.assertNotContains(response, other_user.username)
        self.assertContains(response, "2 time events")

        response = self.client.post(
            reverse("portal:company_approve_dtr", args=(own_log.pk,)),
            {"event": "time_in"},
        )

        self.assertRedirects(response, reverse("portal:company_dtr_queue"))
        own_log.refresh_from_db()
        self.assertTrue(own_log.time_in_approved)
        self.assertFalse(own_log.time_out_approved)
        self.assertFalse(own_log.approved)

        self.client.post(
            reverse("portal:company_approve_dtr", args=(own_log.pk,)),
            {"event": "time_out"},
        )
        own_log.refresh_from_db()
        self.assertTrue(own_log.approved)

    def test_incomplete_attendance_allows_its_recorded_time_to_be_approved(self):
        Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            status=Application.Status.ACCEPTED,
        )
        AttendanceLog.objects.create(
            intern=self.intern,
            work_date="2026-10-01",
            clock_in="2026-10-01T09:00:00Z",
        )
        self.client.force_login(self.company_user)

        response = self.client.get(reverse("portal:company_dtr_queue"))

        self.assertContains(response, "Approve time in")
        self.assertContains(response, "Time out approval")
        self.assertContains(response, "Not recorded")

    def test_company_scorecard_is_limited_to_accepted_interns(self):
        Application.objects.create(
            intern=self.intern,
            posting=self.posting,
            status=Application.Status.ACCEPTED,
        )
        self.client.force_login(self.company_user)

        response = self.client.get(reverse("portal:company_scorecards"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.intern_user.username)

        response = self.client.post(
            reverse("portal:company_scorecards"),
            {
                "intern_id": self.intern.pk,
                "week_start": "2026-10-05",
                "technical": 5,
                "communication": 4,
                "initiative": 3,
                "reliability": 2,
                "feedback": "Strong progress this week.",
            },
        )

        self.assertRedirects(response, reverse("portal:company_scorecards"))
        scorecard = Scorecard.objects.get(company=self.company, intern=self.intern)
        self.assertEqual(scorecard.reviewer, self.company_user)
        self.assertEqual(scorecard.average, 3.5)

        unaccepted_user = User.objects.create_user(username="unaccepted", password="Safe-example-Password-927!")
        unaccepted_intern = InternProfile.objects.get(user=unaccepted_user)
        response = self.client.post(
            reverse("portal:company_scorecards"),
            {
                "intern_id": unaccepted_intern.pk,
                "week_start": "2026-10-05",
                "technical": 5,
                "communication": 5,
                "initiative": 5,
                "reliability": 5,
            },
        )
        self.assertEqual(response.status_code, 404)

    def test_coordinator_can_verify_company_and_review_audit_log(self):
        coordinator = User.objects.create_user(
            username="coordinator",
            password="Safe-example-Password-927!",
            role=User.Role.COORDINATOR,
        )
        self.client.force_login(coordinator)

        response = self.client.get(reverse("portal:admin_verify"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.company.organization)

        response = self.client.post(
            reverse("portal:admin_verify_company", args=(self.company.pk, "verify"))
        )
        self.assertRedirects(response, reverse("portal:admin_verify"))
        self.company.refresh_from_db()
        self.assertTrue(self.company.verified)

        response = self.client.get(reverse("portal:admin_audit"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "company.verified")

    def test_application_unique_constraint_blocks_duplicate(self):
        Application.objects.create(intern=self.intern, posting=self.posting)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Application.objects.create(intern=self.intern, posting=self.posting)

    def test_attendance_unique_constraint_blocks_duplicate_day(self):
        AttendanceLog.objects.create(intern=self.intern, work_date="2026-01-05")
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                AttendanceLog.objects.create(intern=self.intern, work_date="2026-01-05")

    def test_weekly_report_unique_constraint_blocks_duplicate_week(self):
        WeeklyReport.objects.create(intern=self.intern, week_start="2026-01-05", accomplishments="Completed onboarding.")
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                WeeklyReport.objects.create(intern=self.intern, week_start="2026-01-05", accomplishments="Second report.")

    def test_clock_in_uses_server_time_and_saves_enrolled_worksite_coordinates(self):
        fixed_now = datetime(2026, 1, 5, 9, 30, tzinfo=datetime_timezone.utc)
        self._complete_intern_profile()
        self._enroll_attendance_location()
        self.client.force_login(self.intern_user)
        with (
            patch("portal.views.timezone.now", return_value=fixed_now),
            patch(
                "portal.views._attendance_face_check",
                return_value=(AttendanceLog.FaceCheckStatus.MATCHED, 0.82, b"encrypted frame"),
            ),
        ):
            response = self.client.post(
                reverse("portal:attendance"),
                self._attendance_payload(
                    "clock_in",
                    notes="Started inventory review",
                    clock_in="1999-01-01T00:00:00Z",
                ),
            )
        self.assertRedirects(response, reverse("portal:attendance"))
        log = AttendanceLog.objects.get(intern=self.intern)
        self.assertEqual(log.clock_in, fixed_now)
        self.assertEqual(log.notes, "Started inventory review")
        self.assertEqual(str(log.clock_in_latitude), "14.600000")
        self.assertTrue(log.clock_in_within_location)
        self.assertEqual(log.clock_in_face_status, AttendanceLog.FaceCheckStatus.MATCHED)

    def test_clock_out_without_clock_in_does_not_create_empty_record(self):
        self._complete_intern_profile()
        self.client.force_login(self.intern_user)
        response = self.client.post(reverse("portal:attendance"), {"action": "clock_out"})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(AttendanceLog.objects.filter(intern=self.intern).exists())

    def test_coordinator_can_approve_time_in_before_time_out(self):
        coordinator = User.objects.create_user(
            username="coordinator",
            password="Safe-example-Password-927!",
            role=User.Role.COORDINATOR,
        )
        log = AttendanceLog.objects.create(intern=self.intern, work_date="2026-01-05", clock_in="2026-01-05T09:00:00Z")
        self.client.force_login(coordinator)
        response = self.client.post(
            reverse("portal:approve_dtr", args=(log.pk,)),
            {"event": "time_in"},
        )
        self.assertEqual(response.status_code, 302)
        log.refresh_from_db()
        self.assertTrue(log.time_in_approved)
        self.assertFalse(log.time_out_approved)
        self.assertFalse(log.approved)

    def test_invalid_scorecard_does_not_create_placeholder(self):
        coordinator = User.objects.create_user(
            username="coordinator",
            password="Safe-example-Password-927!",
            role=User.Role.COORDINATOR,
        )
        self.client.force_login(coordinator)
        response = self.client.post(
            reverse("portal:scorecards"),
            {
                "intern_id": self.intern.pk,
                "week_start": "2026-01-05",
                "technical": 0,
                "communication": 3,
                "initiative": 3,
                "reliability": 3,
                "feedback": "Review",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Scorecard.objects.count(), 0)

    def test_analytics_shows_saved_risk_assessments(self):
        coordinator = User.objects.create_user(
            username="coordinator",
            password="Safe-example-Password-927!",
            role=User.Role.COORDINATOR,
        )
        RiskAssessment.objects.create(
            intern=self.intern,
            risk_score=72,
            level=RiskAssessment.Level.HIGH,
            indicators=["missed days"],
        )
        self.client.force_login(coordinator)

        response = self.client.get(reverse("portal:analytics"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.intern_user.username)
        self.assertContains(response, "72 / 100")
        self.assertContains(response, "missed days")

    def test_haversine_returns_equatorial_distance(self):
        self.assertAlmostEqual(haversine_distance_km(0, 0, 0, 1), 111.2, delta=0.2)

    def test_upload_validator_checks_pdf_magic_bytes(self):
        validator = FileSizeAndTypeValidator()
        validator(SimpleUploadedFile("resume.pdf", b"%PDF-1.7\ncontent"))
        with self.assertRaises(ValidationError):
            validator(SimpleUploadedFile("fake.pdf", b"plain text disguised as a PDF"))

    def test_application_view_rejects_disguised_upload(self):
        self._complete_intern_profile()
        self.client.force_login(self.intern_user)
        response = self.client.post(
            reverse("portal:apply", args=(self.posting.pk,)),
            {"cover_letter": "Interested", "resume": SimpleUploadedFile("fake.pdf", b"not a real PDF")},
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Application.objects.filter(intern=self.intern, posting=self.posting).exists())

    def test_blank_student_ids_do_not_collide(self):
        self.intern.student_id = ""
        self.intern.save()
        User.objects.create_user(username="intern-two", password="Safe-example-Password-927!")
        self.assertEqual(InternProfile.objects.filter(student_id__isnull=True).count(), 2)